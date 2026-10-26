"""Carry the presenter notes from the 56-slide deck onto a re-paginated deck.
Groups: a slide whose title does not end in '(cont.)' starts a group; '(cont.)' slides join it."""
import json, re, sys
from pptx import Presentation

# usage: remap_notes.py BASELINE.pptx BASELINE_NOTES.json NEW_RAW.pptx OUT.pptx
OLD, NOTES, NEW, OUT = sys.argv[1:5]
# Case-specific note overrides live in notes_overrides.json:
# {"extra": {"<slide title>": ["replace"|"append", "text"]}, "newGroups": {"<slide title>": 101}, "mseTitle": "MSE on admission · <date>", "genogram": "genogram_slide.png"}
import os
OV = json.load(open('notes_overrides.json')) if os.path.exists('notes_overrides.json') else {}

notes = {int(k): v for k, v in json.load(open(NOTES)).items()}

def title(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split('\n')[0]
    return ''
def groups(pres):
    g = []
    for i, s in enumerate(pres.slides, 1):
        t = title(s)
        if t.endswith('(cont.)') and g: g[-1].append(i)
        else: g.append([i])
    return g

old, new = Presentation(OLD), Presentation(NEW)
go, gn = groups(old), groups(new)
def sentences(t):
    return [x for x in re.split(r'(?<=[.!?])\s+(?=[A-Z0-9“\'(])', t) if x]
# Align old and new groups by their first-slide titles; groups new to this deck take notes from NEWG
import difflib
base = lambda t: re.sub(r'\s*\(cont\.\)$', '', t)
to = [base(title(old.slides[g[0]-1])) for g in go]; tn = [base(title(new.slides[g[0]-1])) for g in gn]
NEWG = OV.get("newGroups", {})
pairs = []  # (old group or None, new group)
for op, a1, a2, b1, b2 in difflib.SequenceMatcher(None, to, tn, autojunk=False).get_opcodes():
    if op == 'equal' or (op == 'replace' and a2-a1 == b2-b1):
        pairs += [(go[a1+k], gn[b1+k]) for k in range(b2-b1)]
    elif op == 'insert':
        for k in range(b1, b2):
            assert tn[k] in NEWG, ('unmatched new group', tn[k]); pairs.append(([NEWG[tn[k]]], gn[k]))
    else:
        raise SystemExit(f'cannot align {op} {to[a1:a2]} -> {tn[b1:b2]}')
plan = {}
oldgroup_of = {}
for gi, (a, b) in enumerate(pairs):
    for i in a: oldgroup_of[i] = gi
    sec = sum(notes[i][0] for i in a)
    text = ' '.join(notes[i][1] for i in a if notes[i][1])
    if len(b) == len(a):
        for i, j in zip(a, b): plan[j] = (notes[i][0], notes[i][1])
        continue
    sents = sentences(text); n = len(b)
    per = [sents[k * len(sents) // n:(k + 1) * len(sents) // n] for k in range(n)]
    base_ = sec // n
    for k, j in enumerate(b):
        s_ = base_ + (sec - base_ * n if k == 0 else 0)
        body = ' '.join(per[k]) if per[k] else '(continued)'
        plan[j] = (s_, body)
gn = [b for _, b in pairs]

EXTRA = {k: tuple(v) for k, v in OV.get('extra', {}).items()}
for gi, b in enumerate(gn):
    t = title(new.slides[b[0]-1])
    if t in EXTRA:
        how, txt = EXTRA[t]; sec, body = plan[b[0]]
        plan[b[0]] = (sec, txt if how == "replace" else (body + "\n\n" + txt))
total = 0
for j, s in enumerate(new.slides, 1):
    sec, text = plan[j]; total += sec
    m, r = divmod(total, 60)
    s.notes_slide.notes_text_frame.text = f"[{sec//60}:{sec%60:02d} on this slide · reach {m}:{r:02d} by the end]" + ("\n" + text if text else "")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().lower() == 'mental status examination':
            runs = sh.text_frame.paragraphs[0].runs
            runs[0].text = OV.get('mseTitle', 'Mental status examination')
            for x in runs[1:]: x.text = ''
# GENOGRAM: replace the generated genogram with the hand-drawn figure
from pptx.util import Inches
for sl in new.slides:
    if OV.get('genogram') and title(sl) == 'Family genogram':
        for sh in list(sl.shapes):
            if not sh.is_placeholder:
                sh._element.getparent().remove(sh._element)
        h = 3.95; w = h * 1920 / 1225
        sl.shapes.add_picture(OV['genogram'], Inches((13.333 - w) / 2), Inches(2.55), Inches(w), Inches(h))
new.save(OUT)
# old slide index -> new slide numbers (for the rehearsal file's section labels)
json.dump({str(i): pairs[oldgroup_of[i]][1] for i in oldgroup_of}, open('slide_map.json', 'w'))
print('slides', len(new.slides), 'total', divmod(total, 60))
